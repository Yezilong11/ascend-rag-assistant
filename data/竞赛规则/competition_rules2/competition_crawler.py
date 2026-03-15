"""
大学生竞赛规则爬虫系统
功能：爬取指定大学生竞赛的参赛要求、赛制流程、评审标准等核心信息
输出格式：Markdown、DOCX、JSON
"""

import requests
import asyncio
import aiohttp
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse
import json
import re
from datetime import datetime
import logging
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass, asdict
import concurrent.futures
import time
import os
from pathlib import Path

# 第三方库需要安装：pip install aiohttp beautifulsoup4 markdown python-docx
import markdown
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@dataclass
class CompetitionRule:
    """竞赛规则数据结构"""
    competition_name: str
    competition_url: str
    # 核心信息
    eligibility: str  # 参赛要求
    competition_process: str  # 赛制流程
    judging_criteria: str  # 评审标准
    registration_info: str  # 报名信息
    schedule: str  # 时间安排
    awards: str  # 奖项设置
    # 元数据
    last_updated: str
    source: str
    # 扩展信息
    organizers: str = ""
    level: str = ""  # 国家级/省级等
    categories: str = ""  # 竞赛类别
    official_documents: List[str] = None  # 官方文档链接
    
    def __post_init__(self):
        if self.official_documents is None:
            self.official_documents = []

class CompetitionSpider:
    """竞赛爬虫基类"""
    
    def __init__(self, max_workers: int = 5, timeout: int = 30):
        self.max_workers = max_workers
        self.timeout = timeout
        self.session = None
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
            'Accept-Encoding': 'gzip, deflate, br',
        }
        
    async def __aenter__(self):
        """异步上下文管理器入口"""
        self.session = aiohttp.ClientSession(headers=self.headers)
        return self
        
    async def __aexit__(self, exc_type, exc_val, exc_tb):
        """异步上下文管理器出口"""
        if self.session:
            await self.session.close()
    
    def sync_fetch(self, url: str) -> Optional[str]:
        """同步获取网页内容"""
        try:
            response = requests.get(url, headers=self.headers, timeout=self.timeout)
            response.raise_for_status()
            response.encoding = response.apparent_encoding or 'utf-8'
            return response.text
        except Exception as e:
            logger.error(f"同步获取 {url} 失败: {e}")
            return None
    
    async def async_fetch(self, url: str) -> Optional[str]:
        """异步获取网页内容"""
        try:
            async with self.session.get(url, timeout=self.timeout) as response:
                if response.status == 200:
                    return await response.text(encoding='utf-8')
                else:
                    logger.error(f"异步获取 {url} 失败，状态码: {response.status}")
                    return None
        except Exception as e:
            logger.error(f"异步获取 {url} 失败: {e}")
            return None
    
    def extract_rules_generic(self, html: str, url: str) -> CompetitionRule:
        """通用规则提取方法（子类可重写）"""
        soup = BeautifulSoup(html, 'html.parser')
        
        # 尝试提取标题
        title = soup.find('h1') or soup.find('title')
        competition_name = title.get_text(strip=True) if title else "未知竞赛"
        
        # 通用关键词匹配
        keywords = {
            'eligibility': ['参赛要求', '参赛资格', '报名条件', '参赛对象', '参赛人员'],
            'process': ['赛制', '流程', '比赛流程', '赛程安排', '比赛方式'],
            'judging': ['评审标准', '评分标准', '评审规则', '评分细则', '评选标准'],
            'registration': ['报名', '注册', '报名方式', '报名时间'],
            'schedule': ['时间安排', '赛程', '时间表', '日程安排'],
            'awards': ['奖项', '奖励', '奖项设置', '奖品']
        }
        
        # 提取各部分内容
        extracted_data = {}
        for key, key_list in keywords.items():
            content = []
            for keyword in key_list:
                # 查找包含关键词的元素
                elements = soup.find_all(string=re.compile(keyword))
                for element in elements:
                    # 获取相关段落
                    parent = element.parent
                    if parent.name in ['p', 'div', 'section', 'article']:
                        # 获取后续几个段落
                        for sibling in parent.find_next_siblings(['p', 'div'])[:3]:
                            text = sibling.get_text(strip=True)
                            if text and len(text) > 10:  # 过滤过短文本
                                content.append(text)
            extracted_data[key] = '\n\n'.join(content[:5])  # 限制段落数量
        
        return CompetitionRule(
            competition_name=competition_name,
            competition_url=url,
            eligibility=extracted_data.get('eligibility', ''),
            competition_process=extracted_data.get('process', ''),
            judging_criteria=extracted_data.get('judging', ''),
            registration_info=extracted_data.get('registration', ''),
            schedule=extracted_data.get('schedule', ''),
            awards=extracted_data.get('awards', ''),
            last_updated=datetime.now().strftime('%Y-%m-%d'),
            source=url
        )
    
    async def crawl_competition(self, name: str, url: str) -> Optional[CompetitionRule]:
        """爬取单个竞赛"""
        logger.info(f"开始爬取: {name}")
        
        html = await self.async_fetch(url)
        if not html:
            # 尝试同步方式
            html = self.sync_fetch(url)
            
        if html:
            try:
                rule = self.extract_rules_generic(html, url)
                rule.competition_name = name  # 使用提供的名称
                logger.info(f"成功爬取: {name}")
                return rule
            except Exception as e:
                logger.error(f"解析 {name} 失败: {e}")
                return None
        else:
            logger.warning(f"无法获取 {name} 的页面内容")
            return None
    
    async def crawl_multiple(self, competitions: List[Tuple[str, str]]) -> List[CompetitionRule]:
        """并发爬取多个竞赛"""
        tasks = []
        for name, url in competitions:
            task = self.crawl_competition(name, url)
            tasks.append(task)
        
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        # 过滤有效结果
        valid_results = []
        for result in results:
            if isinstance(result, CompetitionRule):
                valid_results.append(result)
            elif isinstance(result, Exception):
                logger.error(f"任务执行失败: {result}")
        
        return valid_results

class CompetitionParser:
    """竞赛解析器（包含特定竞赛的解析规则）"""
    
    @staticmethod
    def parse_challenge_cup(html: str, url: str) -> CompetitionRule:
        """解析挑战杯竞赛"""
        soup = BeautifulSoup(html, 'html.parser')
        
        # 这里可以添加针对挑战杯网站的具体解析逻辑
        # 示例：查找特定的CSS选择器或HTML结构
        
        rule = CompetitionRule(
            competition_name="挑战杯全国大学生课外学术科技作品竞赛",
            competition_url=url,
            eligibility="",
            competition_process="",
            judging_criteria="",
            registration_info="",
            schedule="",
            awards="",
            last_updated=datetime.now().strftime('%Y-%m-%d'),
            source=url
        )
        
        # 具体解析逻辑...
        return rule
    
    @staticmethod
    def parse_icpc(html: str, url: str) -> CompetitionRule:
        """解析ICPC竞赛"""
        # 类似地实现ICPC特定解析
        pass

class DataExporter:
    """数据导出器"""
    
    @staticmethod
    def to_markdown(rule: CompetitionRule, output_dir: str = "output") -> str:
        """导出为Markdown格式"""
        # 创建输出目录
        os.makedirs(output_dir, exist_ok=True)
        
        # 生成文件名（移除非法字符）
        safe_name = re.sub(r'[\\/*?:"<>|]', '', rule.competition_name)
        filename = f"{safe_name}.md"
        filepath = os.path.join(output_dir, filename)
        
        # 构建Markdown内容
        md_content = f"""# {rule.competition_name}

## 基本信息
- **竞赛名称**: {rule.competition_name}
- **官方网站**: [{rule.competition_url}]({rule.competition_url})
- **最后更新**: {rule.last_updated}
- **数据来源**: {rule.source}

## 参赛要求
{rule.eligibility or '暂无信息'}

## 赛制流程
{rule.competition_process or '暂无信息'}

## 评审标准
{rule.judging_criteria or '暂无信息'}

## 报名信息
{rule.registration_info or '暂无信息'}

## 时间安排
{rule.schedule or '暂无信息'}

## 奖项设置
{rule.awards or '暂无信息'}

## 主办单位
{rule.organizers or '暂无信息'}

## 竞赛级别
{rule.level or '暂无信息'}

## 竞赛类别
{rule.categories or '暂无信息'}

## 官方文档
{chr(10).join(f"- {doc}" for doc in rule.official_documents) if rule.official_documents else '暂无'}
"""
        
        # 保存文件
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(md_content)
        
        logger.info(f"已保存Markdown文件: {filepath}")
        return filepath
    
    @staticmethod
    def to_docx(rule: CompetitionRule, output_dir: str = "output") -> str:
        """导出为DOCX格式"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            # 创建输出目录
            os.makedirs(output_dir, exist_ok=True)
            
            # 生成文件名
            safe_name = re.sub(r'[\\/*?:"<>|]', '', rule.competition_name)
            filename = f"{safe_name}.docx"
            filepath = os.path.join(output_dir, filename)
            
            # 创建文档
            doc = Document()
            
            # 添加标题
            title = doc.add_heading(rule.competition_name, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # 添加基本信息表格
            doc.add_heading('基本信息', level=1)
            table = doc.add_table(rows=4, cols=2)
            table.style = 'Light Grid Accent 1'
            
            # 填充表格
            data = [
                ('竞赛名称', rule.competition_name),
                ('官方网站', rule.competition_url),
                ('最后更新', rule.last_updated),
                ('数据来源', rule.source)
            ]
            
            for i, (key, value) in enumerate(data):
                table.cell(i, 0).text = key
                table.cell(i, 1).text = value
            
            # 添加各个部分
            sections = [
                ('参赛要求', rule.eligibility),
                ('赛制流程', rule.competition_process),
                ('评审标准', rule.judging_criteria),
                ('报名信息', rule.registration_info),
                ('时间安排', rule.schedule),
                ('奖项设置', rule.awards),
                ('主办单位', rule.organizers),
                ('竞赛级别', rule.level),
                ('竞赛类别', rule.categories)
            ]
            
            for title, content in sections:
                if content:
                    doc.add_heading(title, level=2)
                    doc.add_paragraph(content)
            
            # 保存文档
            doc.save(filepath)
            logger.info(f"已保存DOCX文件: {filepath}")
            return filepath
            
        except ImportError:
            logger.warning("未安装python-docx库，无法导出DOCX格式")
            return ""
    
    @staticmethod
    def to_json(rules: List[CompetitionRule], output_dir: str = "output") -> str:
        """导出为JSON格式"""
        os.makedirs(output_dir, exist_ok=True)
        
        # 转换为字典列表
        data = [asdict(rule) for rule in rules]
        
        filepath = os.path.join(output_dir, "competition_rules.json")
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"已保存JSON文件: {filepath}")
        return filepath
    
    @staticmethod
    def export_all(rules: List[CompetitionRule], output_dir: str = "output"):
        """导出所有格式"""
        # 导出JSON
        DataExporter.to_json(rules, output_dir)
        
        # 导出单个文件的Markdown和DOCX
        md_files = []
        docx_files = []
        
        for rule in rules:
            md_file = DataExporter.to_markdown(rule, os.path.join(output_dir, "markdown"))
            if md_file:
                md_files.append(md_file)
            
            docx_file = DataExporter.to_docx(rule, os.path.join(output_dir, "docx"))
            if docx_file:
                docx_files.append(docx_file)
        
        # 创建汇总文件
        DataExporter.create_summary(rules, output_dir)
        
        return {
            'markdown': md_files,
            'docx': docx_files,
            'json': os.path.join(output_dir, "competition_rules.json")
        }
    
    @staticmethod
    def create_summary(rules: List[CompetitionRule], output_dir: str):
        """创建汇总报告"""
        summary_md = f"""# 大学生竞赛规则汇总报告

## 统计信息
- **爬取时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
- **竞赛总数**: {len(rules)}
- **成功爬取**: {len([r for r in rules if any([r.eligibility, r.competition_process, r.judging_criteria])])}

## 竞赛列表
"""
        
        for i, rule in enumerate(rules, 1):
            has_data = any([rule.eligibility, rule.competition_process, rule.judging_criteria])
            status = "✅" if has_data else "⚠️"
            
            summary_md += f"""
### {i}. {status} {rule.competition_name}

**官方网站**: [{rule.competition_url}]({rule.competition_url})

**数据完整性**:
- 参赛要求: {'✅' if rule.eligibility else '❌'}
- 赛制流程: {'✅' if rule.competition_process else '❌'}
- 评审标准: {'✅' if rule.judging_criteria else '❌'}

---
"""
        
        # 保存汇总文件
        summary_path = os.path.join(output_dir, "SUMMARY.md")
        with open(summary_path, 'w', encoding='utf-8') as f:
            f.write(summary_md)
        
        logger.info(f"已保存汇总文件: {summary_path}")

class CompetitionManager:
    """竞赛管理器"""
    
    def __init__(self):
        self.competitions = []
        self.load_default_competitions()
    
    def load_default_competitions(self):
        """加载默认竞赛列表（基于用户提供的列表）"""
        # 这里只添加部分示例，实际使用时需要补充完整的URL
        default_competitions = [
            ("中国国际大学生创新大赛", "http://www.moe.gov.cn/"),
            ("挑战杯全国大学生课外学术科技作品竞赛", "http://www.tiaozhanbei.net/"),
            ("挑战杯中国大学生创业计划大赛", "http://www.chuangqingbei.net/"),
            ("全国大学生数学建模竞赛", "http://www.mcm.edu.cn/"),
            ("全国大学生电子设计竞赛", "http://www.nuedc.com.cn/"),
            ("全国大学生智能汽车竞赛", "http://www.smartcar.org.cn/"),
            ("全国大学生英语竞赛", "http://www.chinaneccs.cn/"),
            ("全国大学生广告艺术大赛", "http://www.sun-ada.net/"),
            ("全国大学生计算机设计大赛", "http://jsjds.blcu.edu.cn/"),
            ("国际大学生程序设计竞赛", "https://icpc.global/"),
        ]
        self.competitions = default_competitions
    
    def add_competition(self, name: str, url: str):
        """添加竞赛"""
        self.competitions.append((name, url))
    
    def load_from_file(self, filepath: str):
        """从文件加载竞赛列表"""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith('#'):
                        parts = line.split(',')
                        if len(parts) >= 2:
                            name = parts[0].strip()
                            url = parts[1].strip()
                            self.competitions.append((name, url))
        except FileNotFoundError:
            logger.error(f"文件不存在: {filepath}")
    
    def get_competitions(self) -> List[Tuple[str, str]]:
        """获取竞赛列表"""
        return self.competitions

async def main():
    """主函数"""
    print("=" * 60)
    print("大学生竞赛规则爬虫系统")
    print("=" * 60)
    
    # 初始化管理器
    manager = CompetitionManager()
    
    # 用户可以选择从文件加载
    config_file = "competitions.csv"
    if os.path.exists(config_file):
        print(f"检测到配置文件: {config_file}")
        manager.load_from_file(config_file)
    
    # 显示要爬取的竞赛
    competitions = manager.get_competitions()
    print(f"\n准备爬取 {len(competitions)} 个竞赛:")
    for i, (name, url) in enumerate(competitions, 1):
        print(f"{i:2d}. {name}")
    
    # 确认开始
    input("\n按Enter键开始爬取...")
    
    # 创建爬虫实例
    async with CompetitionSpider(max_workers=10) as spider:
        # 开始爬取
        print("\n开始爬取竞赛规则...")
        start_time = time.time()
        
        rules = await spider.crawl_multiple(competitions)
        
        elapsed_time = time.time() - start_time
        print(f"\n爬取完成！耗时: {elapsed_time:.2f}秒")
        print(f"成功爬取: {len(rules)}/{len(competitions)} 个竞赛")
        
        # 导出数据
        if rules:
            print("\n正在导出数据...")
            exporter = DataExporter()
            results = exporter.export_all(rules, "competition_rules_output")
            
            print("\n导出完成！")
            print(f"JSON文件: {results['json']}")
            print(f"Markdown文件: {len(results['markdown'])} 个")
            print(f"DOCX文件: {len(results['docx'])} 个")
            print(f"汇总报告: competition_rules_output/SUMMARY.md")
        else:
            print("未成功爬取到任何数据")
    
    print("\n程序执行完毕！")

def create_sample_config():
    """创建示例配置文件"""
    sample_content = """# 大学生竞赛列表配置文件
# 格式：竞赛名称,官方网站URL

# A类竞赛
中国国际大学生创新大赛,http://www.moe.gov.cn/
挑战杯全国大学生课外学术科技作品竞赛,http://www.tiaozhanbei.net/
全国大学生数学建模竞赛,http://www.mcm.edu.cn/
全国大学生电子设计竞赛,http://www.nuedc.com.cn/

# B类竞赛
全国大学生智能汽车竞赛,http://www.smartcar.org.cn/
全国大学生英语竞赛,http://www.chinaneccs.cn/
全国大学生广告艺术大赛,http://www.sun-ada.net/

# 国际竞赛
国际大学生程序设计竞赛,https://icpc.global/
美国大学生数学建模竞赛,https://www.comap.com/
"""
    
    with open("competitions_sample.csv", "w", encoding="utf-8") as f:
        f.write(sample_content)
    print("已创建示例配置文件: competitions_sample.csv")

if __name__ == "__main__":
    # 创建示例配置文件
    create_sample_config()
    
    # 运行主程序
    asyncio.run(main())