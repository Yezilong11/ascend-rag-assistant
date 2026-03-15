"""
竞赛技术文档爬虫
功能：爬取Word文档中所有竞赛官网的技术文档
作者：AI Assistant
"""

import re
import os
import json
import time
import requests
from urllib.parse import urljoin, urlparse
from bs4 import BeautifulSoup
from docx import Document
from pathlib import Path
from datetime import datetime
import html2text

class CompetitionDocCrawler:
    def __init__(self, word_file_path, output_dir="competition_docs"):
        self.word_file_path = word_file_path
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
        
        # 创建子目录
        self.markdown_dir = self.output_dir / "markdown"
        self.raw_dir = self.output_dir / "raw_html"
        self.metadata_file = self.output_dir / "metadata.json"
        
        self.markdown_dir.mkdir(exist_ok=True)
        self.raw_dir.mkdir(exist_ok=True)
        
        # 初始化HTML转Markdown转换器
        self.html2text_converter = html2text.HTML2Text()
        self.html2text_converter.ignore_links = False
        self.html2text_converter.ignore_images = False
        self.html2text_converter.body_width = 0  # 不自动换行
        
        # 请求配置
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        }
        self.session = requests.Session()
        self.session.headers.update(self.headers)
        
        # 存储结果
        self.competitions = []
        self.failed_urls = []
        
    def extract_links_from_word(self):
        """从Word文档提取竞赛名称和链接"""
        print(f"📄 正在解析Word文档: {self.word_file_path}")
        
        doc = Document(self.word_file_path)
        content = []
        
        # 提取所有段落文本
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content.append(text)
        
        # 提取表格中的内容
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        content.append(text)
        
        full_text = '\n'.join(content)
        
        # 使用正则表达式匹配竞赛名称和URL
        # 匹配模式：**竞赛名称**：[URL] 或 数字. **竞赛名称**：[URL]
        pattern = r'(?:\d+\.\s*)?\*\*(.+?)\*\*.*?[：:]\s*\[?(https?://[^\]\s]+)\]?'
        matches = re.findall(pattern, full_text)
        
        # 如果没匹配到，尝试其他模式
        if not matches:
            # 尝试匹配纯URL
            url_pattern = r'(https?://[^\s\]\)]+)'
            urls = re.findall(url_pattern, full_text)
            
            # 为每个URL尝试找到对应的竞赛名称
            lines = full_text.split('\n')
            for url in urls:
                for line in lines:
                    if url in line:
                        # 提取可能的竞赛名称（通常在** **之间或在"竞赛"附近）
                        name_match = re.search(r'\*\*(.+?)\*\*', line)
                        if name_match:
                            matches.append((name_match.group(1), url))
                            break
        
        # 去重并整理
        seen = set()
        for name, url in matches:
            # 清理URL（移除markdown格式）
            clean_url = url.rstrip(')').rstrip(']').rstrip('.')
            # 清理名称
            clean_name = name.strip()
            
            if clean_url not in seen and clean_url.startswith('http'):
                seen.add(clean_url)
                self.competitions.append({
                    'name': clean_name,
                    'url': clean_url,
                    'status': 'pending',
                    'docs_found': []
                })
        
        print(f"✅ 成功提取 {len(self.competitions)} 个竞赛链接")
        for i, comp in enumerate(self.competitions[:5], 1):
            print(f"   {i}. {comp['name'][:30]}... -> {comp['url'][:50]}...")
        
        if len(self.competitions) > 5:
            print(f"   ... 还有 {len(self.competitions)-5} 个")
            
        return self.competitions
    
    def find_documentation_links(self, url, depth=0):
        """在网页中查找技术文档链接"""
        if depth > 2:  # 限制递归深度
            return []
        
        try:
            print(f"   🔍 正在分析页面: {url[:60]}...")
            response = self.session.get(url, timeout=15, allow_redirects=True)
            response.encoding = response.apparent_encoding or 'utf-8'
            
            soup = BeautifulSoup(response.text, 'html.parser')
            doc_links = []
            
            # 关键词匹配（技术文档相关）
            keywords = [
                '技术文档', '文档中心', '开发文档', '比赛规则', '竞赛指南', 
                '参赛手册', '技术规范', '赛题说明', '资料下载', '通知公告',
                'documentation', 'docs', 'guide', 'manual', 'rules', 
                'download', 'resources', 'technical', 'specification',
                '赛题', '规则', '通知', '下载', '资料', '指南'
            ]
            
            # 1. 查找导航菜单中的文档链接
            nav_selectors = ['nav', 'header', '.menu', '.navbar', '.navigation', 
                           '.sidebar', '.nav', '#menu', '#nav']
            for selector in nav_selectors:
                for elem in soup.select(selector):
                    links = elem.find_all('a', href=True)
                    for link in links:
                        text = link.get_text(strip=True).lower()
                        href = link['href']
                        if any(kw in text for kw in keywords):
                            full_url = urljoin(url, href)
                            doc_links.append({
                                'url': full_url,
                                'text': link.get_text(strip=True),
                                'type': 'navigation',
                                'source': url
                            })
            
            # 2. 查找主要内容区域的链接
            content_selectors = ['main', '.content', '.main', '#content', 
                               'article', '.container', '.wrapper']
            for selector in content_selectors:
                for elem in soup.select(selector):
                    links = elem.find_all('a', href=True)
                    for link in links:
                        text = link.get_text(strip=True)
                        href = link['href']
                        text_lower = text.lower()
                        
                        # 匹配关键词
                        if any(kw in text_lower for kw in keywords):
                            full_url = urljoin(url, href)
                            # 排除外部链接
                            if self.is_same_domain(url, full_url):
                                doc_links.append({
                                    'url': full_url,
                                    'text': text,
                                    'type': 'content',
                                    'source': url
                                })
            
            # 3. 查找特定的文档页面模式
            doc_patterns = [
                r'/(doc|docs|document|documentation|guide|manual|rule|notice|download|resource)',
                r'/(about|intro|competition|contest|match)',
                r'\.(pdf|doc|docx|md|txt)$'
            ]
            
            for link in soup.find_all('a', href=True):
                href = link['href']
                for pattern in doc_patterns:
                    if re.search(pattern, href, re.I):
                        full_url = urljoin(url, href)
                        if self.is_same_domain(url, full_url):
                            text = link.get_text(strip=True) or '未命名链接'
                            # 检查是否已存在
                            if not any(d['url'] == full_url for d in doc_links):
                                doc_links.append({
                                    'url': full_url,
                                    'text': text[:50],
                                    'type': 'pattern_match',
                                    'source': url
                                })
            
            # 去重
            seen_urls = set()
            unique_links = []
            for link in doc_links:
                if link['url'] not in seen_urls:
                    seen_urls.add(link['url'])
                    unique_links.append(link)
            
            return unique_links
            
        except Exception as e:
            print(f"   ❌ 获取页面失败: {str(e)}")
            return []
    
    def is_same_domain(self, url1, url2):
        """检查两个URL是否在同一域名下"""
        try:
            domain1 = urlparse(url1).netloc
            domain2 = urlparse(url2).netloc
            # 允许子域名
            return domain1 in domain2 or domain2 in domain1
        except:
            return False
    
    def crawl_page(self, url, competition_name):
        """爬取单个页面并保存"""
        try:
            response = self.session.get(url, timeout=20)
            response.encoding = response.apparent_encoding or 'utf-8'
            
            # 生成文件名
            parsed = urlparse(url)
            path_parts = parsed.path.strip('/').replace('/', '_')
            if not path_parts:
                path_parts = 'index'
            
            timestamp = datetime.now().strftime("%Y%m%d")
            safe_name = re.sub(r'[^\w\s-]', '', competition_name)[:30].strip()
            filename_base = f"{safe_name}_{path_parts}_{timestamp}"
            
            # 保存原始HTML
            html_filename = self.raw_dir / f"{filename_base}.html"
            with open(html_filename, 'w', encoding='utf-8') as f:
                f.write(response.text)
            
            # 转换为Markdown
            markdown_content = self.html2text_converter.handle(response.text)
            
            # 添加元数据头
            header = f"""---
title: {competition_name} - 技术文档
source_url: {url}
crawled_at: {datetime.now().isoformat()}
competition: {competition_name}
---

# {competition_name}

> 原始链接: [{url}]({url})
> 爬取时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

---

"""
            full_markdown = header + markdown_content
            
            # 保存Markdown
            md_filename = self.markdown_dir / f"{filename_base}.md"
            with open(md_filename, 'w', encoding='utf-8') as f:
                f.write(full_markdown)
            
            return {
                'success': True,
                'html_file': str(html_filename),
                'markdown_file': str(md_filename),
                'url': url,
                'size': len(response.text)
            }
            
        except Exception as e:
            return {
                'success': False,
                'url': url,
                'error': str(e)
            }
    
    def crawl_competition(self, competition):
        """爬取单个竞赛的所有文档"""
        name = competition['name']
        url = competition['url']
        
        print(f"\n{'='*60}")
        print(f"🏆 正在处理: {name}")
        print(f"🌐 官网: {url}")
        print(f"{'='*60}")
        
        results = []
        
        # 1. 首先爬取主页
        print("📥 正在爬取主页...")
        home_result = self.crawl_page(url, name)
        results.append(home_result)
        
        if not home_result['success']:
            print(f"❌ 主页爬取失败: {home_result.get('error')}")
            competition['status'] = 'failed'
            self.failed_urls.append({'name': name, 'url': url, 'error': home_result.get('error')})
            return results
        
        print(f"✅ 主页已保存")
        
        # 2. 查找文档链接
        print("🔍 正在查找技术文档链接...")
        doc_links = self.find_documentation_links(url)
        
        # 过滤掉主页
        doc_links = [link for link in doc_links if link['url'] != url]
        
        print(f"📋 发现 {len(doc_links)} 个潜在文档页面")
        
        # 3. 爬取文档页面（限制数量避免过多）
        max_docs = 5  # 每个竞赛最多爬取5个文档页面
        for i, link in enumerate(doc_links[:max_docs], 1):
            print(f"\n   [{i}/{min(len(doc_links), max_docs)}] {link['text'][:40]}...")
            print(f"   URL: {link['url'][:60]}...")
            
            # 检查是否已爬取
            if any(r['url'] == link['url'] for r in results):
                print("   ⏭️  已爬取，跳过")
                continue
            
            result = self.crawl_page(link['url'], name)
            results.append(result)
            
            if result['success']:
                print(f"   ✅ 已保存: {result['markdown_file']}")
                competition['docs_found'].append({
                    'title': link['text'],
                    'url': link['url'],
                    'file': result['markdown_file']
                })
            else:
                print(f"   ❌ 失败: {result.get('error')}")
            
            # 礼貌延迟
            time.sleep(1)
        
        competition['status'] = 'completed'
        competition['total_pages'] = len(results)
        competition['successful_pages'] = sum(1 for r in results if r['success'])
        
        return results
    
    def generate_summary(self):
        """生成汇总文档"""
        summary_md = f"""# 竞赛技术文档爬取汇总

> 生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

## 统计信息

- **总竞赛数**: {len(self.competitions)}
- **成功爬取**: {sum(1 for c in self.competitions if c['status'] == 'completed')}
- **失败**: {sum(1 for c in self.competitions if c['status'] == 'failed')}
- **待处理**: {sum(1 for c in self.competitions if c['status'] == 'pending')}

## 竞赛清单

| 序号 | 竞赛名称 | 官网 | 状态 | 文档数 |
|------|---------|------|------|--------|
"""
        
        for i, comp in enumerate(self.competitions, 1):
            status_emoji = {
                'completed': '✅',
                'failed': '❌',
                'pending': '⏳'
            }.get(comp['status'], '❓')
            
            doc_count = len(comp.get('docs_found', []))
            summary_md += f"| {i} | {comp['name']} | [链接]({comp['url']}) | {status_emoji} | {doc_count} |\n"
        
        summary_md += f"""
## 详细文档索引

"""
        
        for comp in self.competitions:
            if comp.get('docs_found'):
                summary_md += f"\n### {comp['name']}\n\n"
                summary_md += f"- **官网**: {comp['url']}\n"
                summary_md += f"- **状态**: {comp['status']}\n\n"
                
                for doc in comp['docs_found']:
                    # 获取相对路径
                    rel_path = os.path.relpath(doc['file'], self.output_dir)
                    summary_md += f"- [{doc['title']}]({rel_path}) - [原始链接]({doc['url']})\n"
        
        if self.failed_urls:
            summary_md += f"\n## 失败的URL\n\n"
            for item in self.failed_urls:
                summary_md += f"- **{item['name']}**: {item['url']} - 错误: {item['error']}\n"
        
        # 保存汇总文档
        summary_file = self.output_dir / "README.md"
        with open(summary_file, 'w', encoding='utf-8') as f:
            f.write(summary_md)
        
        print(f"\n📊 汇总文档已生成: {summary_file}")
        return summary_file
    
    def save_metadata(self):
        """保存元数据JSON"""
        metadata = {
            'crawl_time': datetime.now().isoformat(),
            'total_competitions': len(self.competitions),
            'output_directory': str(self.output_dir),
            'competitions': self.competitions,
            'failed_urls': self.failed_urls
        }
        
        with open(self.metadata_file, 'w', encoding='utf-8') as f:
            json.dump(metadata, f, ensure_ascii=False, indent=2)
        
        print(f"💾 元数据已保存: {self.metadata_file}")
    
    def run(self):
        """运行完整爬取流程"""
        print("🚀 开始爬取竞赛技术文档...")
        print(f"📁 输出目录: {self.output_dir.absolute()}")
        
        # 1. 提取链接
        self.extract_links_from_word()
        
        if not self.competitions:
            print("❌ 未找到任何竞赛链接，请检查Word文档格式")
            return
        
        # 2. 爬取每个竞赛
        for i, competition in enumerate(self.competitions, 1):
            print(f"\n\n📌 进度: [{i}/{len(self.competitions)}]")
            try:
                self.crawl_competition(competition)
            except Exception as e:
                print(f"❌ 处理竞赛时出错: {str(e)}")
                competition['status'] = 'failed'
                self.failed_urls.append({
                    'name': competition['name'],
                    'url': competition['url'],
                    'error': str(e)
                })
            
            # 每5个竞赛保存一次进度
            if i % 5 == 0:
                self.save_metadata()
                print(f"💾 进度已保存 ({i}/{len(self.competitions)})")
            
            # 礼貌延迟
            time.sleep(2)
        
        # 3. 生成汇总
        self.generate_summary()
        self.save_metadata()
        
        print(f"\n{'='*60}")
        print("✅ 爬取完成!")
        print(f"📁 输出目录: {self.output_dir.absolute()}")
        print(f"📝 Markdown文件: {self.markdown_dir}")
        print(f"🌐 HTML文件: {self.raw_dir}")
        print(f"📊 汇总文档: {self.output_dir / 'README.md'}")
        print(f"{'='*60}")


def main():
    # 配置
    WORD_FILE = "大赛官网（初版）.docx"  # 你的Word文件名
    OUTPUT_DIR = "competition_docs"      # 输出目录
    
    # 检查文件是否存在
    if not os.path.exists(WORD_FILE):
        print(f"❌ 文件不存在: {WORD_FILE}")
        print("请确保Word文件与脚本在同一目录，或修改路径")
        return
    
    # 创建爬虫实例并运行
    crawler = CompetitionDocCrawler(WORD_FILE, OUTPUT_DIR)
    crawler.run()


if __name__ == "__main__":
    main()